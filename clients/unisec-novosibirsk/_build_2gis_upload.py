from __future__ import annotations

import argparse
import csv
from pathlib import Path


BASE_DIR = Path(__file__).resolve().parent
SOURCE_PATH = BASE_DIR / "yandex-maps-services-working.csv"
IMAGE_MAP_PATH = BASE_DIR / "yandex-maps-image-filenames.csv"
CSV_OUTPUT_PATH = BASE_DIR / "2gis-price-list-ready.csv"
XLSX_OUTPUT_PATH = BASE_DIR / "2gis-price-list-ready.xlsx"

IMAGE_BASE_URL = "https://cs836180.tw1.ru/yandex-maps-services/unisec"

# Exact column order from the official 2GIS CSV example.
FIELDS = [
    "name",
    "price",
    "currencyId",
    "measure",
    "available",
    "category",
    "url",
    "picture",
    "vendor",
    "id",
    "description",
]

MEASURE_MAP = {
    "штука": "шт.",
    "пара": "пара",
    "комплект": "комплект",
    "килограмм": "кг",
    "кв. метр": "кв.м",
    "погонный метр": "пог.м",
}


def cp1251_safe(value: str) -> str:
    """Keep the official sample encoding without losing service meaning."""
    replacements = {
        "₽": "руб.",
        "×": "x",
        "²": "2",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = value.replace("руб..", "руб.")
    value.encode("cp1251")
    return value


def read_rows() -> list[dict[str, str]]:
    with SOURCE_PATH.open("r", encoding="utf-8-sig", newline="") as source:
        rows = [row for row in csv.DictReader(source) if row["status"] == "publish"]

    with IMAGE_MAP_PATH.open("r", encoding="utf-8-sig", newline="") as source:
        image_by_keyword = {
            row["primary_keyword"]: row["image_filename"]
            for row in csv.DictReader(source)
        }

    output: list[dict[str, str]] = []
    for row in rows:
        keyword = row["primary_keyword"]
        filename = image_by_keyword.get(keyword)
        if not filename:
            raise ValueError(f"No image filename for: {keyword}")

        measure = MEASURE_MAP.get(row["unit"])
        if not measure:
            raise ValueError(f"Unknown measure: {row['unit']}")

        output.append(
            {
                "name": cp1251_safe(row["name"]),
                "price": row["price"],
                "currencyId": "RUB",
                "measure": measure,
                "available": "В наличии",
                "category": cp1251_safe(row["category"]),
                "url": "",
                "picture": f"{IMAGE_BASE_URL}/{filename}",
                "vendor": "UNISEC",
                "id": Path(filename).stem,
                "description": cp1251_safe(row["description"]),
            }
        )

    if len(output) != 43:
        raise ValueError(f"Expected 43 published services, got {len(output)}")
    if len({row["id"] for row in output}) != len(output):
        raise ValueError("2GIS IDs must be unique")
    return output


def write_csv(rows: list[dict[str, str]]) -> None:
    # The official 2GIS example is semicolon-delimited and encoded in Windows-1251.
    with CSV_OUTPUT_PATH.open("w", encoding="cp1251", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=FIELDS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict[str, str]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "XLSX generation requires openpyxl; the CSV file is still ready for 2GIS"
        ) from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товары и услуги"
    sheet.append(FIELDS)
    for row in rows:
        sheet.append([row[field] for field in FIELDS])

    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [36, 11, 12, 13, 15, 28, 16, 60, 14, 38, 90]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(XLSX_OUTPUT_PATH)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--xlsx",
        action="store_true",
        help="Also build XLSX (requires openpyxl)",
    )
    args = parser.parse_args()

    rows = read_rows()
    write_csv(rows)
    print(f"Created {CSV_OUTPUT_PATH.name}: {len(rows)} services")
    if args.xlsx:
        write_xlsx(rows)
        print(f"Created {XLSX_OUTPUT_PATH.name}: {len(rows)} services")


if __name__ == "__main__":
    main()
