from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
WORKING = ROOT / "yandex-maps-services-working.csv"
FILENAMES = ROOT / "yandex-maps-image-filenames.csv"
READY_CSV = ROOT / "yandex-maps-price-list-ready.csv"
READY_XLSX = ROOT / "yandex-maps-price-list-ready.xlsx"
PHOTO_BASE_URL = "https://cs836180.tw1.ru/yandex-maps-services/unisec"
FIELDS = [
    "Категория",
    "Название",
    "Цена",
    "Количество",
    "Единицы измерения",
    "Описание",
    "Фотография",
    "Популярный товар",
]
POPULAR_KEYWORDS = {
    "химчистка одежды",
    "химчистка обуви",
    "прачечная",
    "ремонт одежды",
    "химчистка штор",
    "чистка подушек",
    "химчистка свадебного платья",
    "реставрация обуви",
    "химчистка одеяла",
    "покраска одежды",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def main() -> None:
    working = [row for row in read_csv(WORKING) if row["status"] == "publish"]
    filenames = {
        row["primary_keyword"].casefold(): row["image_filename"]
        for row in read_csv(FILENAMES)
    }

    rows: list[dict[str, str]] = []
    for row in working:
        keyword = row["primary_keyword"].casefold()
        if keyword not in filenames:
            raise AssertionError(f"No photo filename for {row['primary_keyword']}")
        rows.append(
            {
                "Категория": row["category"],
                "Название": row["name"],
                "Цена": row["price"],
                "Количество": row["quantity"],
                "Единицы измерения": row["unit"],
                "Описание": row["description"],
                "Фотография": f"{PHOTO_BASE_URL}/{filenames[keyword]}",
                "Популярный товар": "Да" if keyword in POPULAR_KEYWORDS else "",
            }
        )

    if len(rows) != 43:
        raise AssertionError(f"Expected 43 publish rows, got {len(rows)}")

    with READY_CSV.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товары и услуги"
    sheet.append(FIELDS)
    for row in rows:
        sheet.append([row[field] for field in FIELDS])

    header_fill = PatternFill("solid", fgColor="0090A7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {"A": 28, "B": 42, "C": 12, "D": 12, "E": 22, "F": 80, "G": 92, "H": 20}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(READY_XLSX)

    print(f"rows={len(rows)}")
    print(READY_CSV)
    print(READY_XLSX)


if __name__ == "__main__":
    main()
