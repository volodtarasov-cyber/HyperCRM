"""Черновики статей сайта Arcars → Word-файлы для согласования с клиентом.

Берёт markdown-статьи SEO-завода из arcars-site/src/content/blog (frontmatter +
тело) и собирает по .docx на статью: шапка «что это и как править», паспорт
страницы (адрес, запрос, title, description), затем сам текст.

Клиент правит в режиме рецензирования или пишет комментарии; правки потом
вносятся обратно в .md руками.

    .tools/venv/Scripts/python.exe .tools/projects/arcars/drafts_to_docx.py <slug> [<slug> ...] [--out DIR] [--notes notes.yml]

notes.yml — замечания редактора по статьям: {slug: ["что проверить", ...]}. Они
встают в начало документа блоком «Что проверить и дописать».

Без --out файлы кладутся в clients/arcars/reports/drafts-YYYY-MM-DD/.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BLOG = Path(r"D:\SecondBrain\arcars-site\src\content\blog")
VAULT = Path(__file__).resolve().parents[3]
SITE = "https://ar-cars.ru"
GREY = RGBColor(0x6B, 0x6F, 0x76)
ACCENT = RGBColor(0xE8, 0x5D, 0x2E)


def split_frontmatter(text: str) -> tuple[dict, str]:
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return {}, text
    return yaml.safe_load(m.group(1)) or {}, text[m.end():]


# --- inline: **жирный**, *курсив*, [текст](адрес), `код`, HTML-теги -------------

INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*\n]+\*(?!\*)|`[^`]+`)")


def add_hyperlink(par, text: str, url: str) -> None:
    if url.startswith("/"):
        url = SITE + url
    r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1F5FBF"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(par, text: str, bold: bool = False) -> None:
    text = re.sub(r"<br\s*/?>", " ", text)
    text = re.sub(r"</?[a-zA-Z][^>]*>", "", text)          # HTML-обёртки завода
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(par, label, url.split()[0])
        elif part.startswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            r = par.add_run(part)
            r.bold = bold or None


# --- блоки ------------------------------------------------------------------

def shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc, rows: list[list[str]]) -> None:
    head, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ""
        add_inline(c.paragraphs[0], h, bold=True); shade(c, "F2F2F2")
    for row in body:
        cells = t.add_row().cells
        for i in range(len(head)):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "")
    doc.add_paragraph()


def note(doc, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.color.rgb = GREY


def render_body(doc, body: str) -> int:
    """Тело статьи в документ. Возвращает число слов текста."""
    lines = body.splitlines()
    words = 0
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1; continue
        if s.startswith("<!--"):
            if "cta" in s.lower():
                note(doc, "[Здесь на сайте будет форма заявки]")
            while "-->" not in lines[i]:
                i += 1
            i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells); words += sum(len(c.split()) for c in cells)
                i += 1
            add_table(doc, rows); continue
        m = re.match(r"^(#{2,4})\s+(.*)", s)
        if m:
            level = len(m.group(1)) - 1
            h = doc.add_heading(level=level); add_inline(h, m.group(2)); i += 1; continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip()); i += 1
            p = doc.add_paragraph(style="Intense Quote"); add_inline(p, " ".join(buf))
            words += len(" ".join(buf).split()); continue
        m = re.match(r"^([-*]|\d+[.)])\s+(.*)", s)
        if m:
            # Нумерованные — номером в тексте: у стиля List Number в python-docx
            # нумерация сквозная, второй список продолжил бы 4, 5, 6.
            if m.group(1)[0].isdigit():
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.63)
                p.add_run(m.group(1).rstrip(".)") + ". ")
            else:
                p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(2))
            words += len(m.group(2).split()); i += 1; continue
        if s.startswith("![") :
            note(doc, "[Иллюстрация]"); i += 1; continue
        buf = [s]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||>|[-*]\s|\d+[.)]\s|<!--)", lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        text = " ".join(buf)
        p = doc.add_paragraph(); add_inline(p, text)
        words += len(re.sub(r"<[^>]+>", "", text).split())
    return words


def build(slug: str, out_dir: Path, notes: list[str] | None = None) -> Path:
    src = BLOG / f"{slug}.md"
    if not src.exists():
        src = BLOG / f"{slug}.mdx"
    fm, body = split_frontmatter(src.read_text(encoding="utf-8"))

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    p = doc.add_paragraph()
    r = p.add_run("Черновик статьи для сайта ar-cars.ru — на согласование"); r.bold = True; r.font.color.rgb = ACCENT
    note(doc, "Проверьте прежде всего факты: цены, сроки, порядок работы, что ARCars делает и чего не делает. "
              "Править можно прямо в тексте в режиме рецензирования (Рецензирование → Исправления) "
              "или оставлять комментарии на полях. Ничего не опубликуем, пока текст не согласован.")

    if notes:
        h = doc.add_paragraph(); r = h.add_run("Что проверить и дописать"); r.bold = True
        for n in notes:
            add_inline(doc.add_paragraph(style="List Bullet"), n)
        doc.add_paragraph()

    doc.add_heading(fm.get("h1") or fm.get("title") or slug, level=0)

    info = [
        ("Адрес на сайте", f"{SITE}/blog/{fm.get('slug', slug)}/"),
        ("Главный поисковый запрос", fm.get("primaryQuery") or (fm.get("seoKeywords") or [""])[0]),
        ("Заголовок в поиске (title)", fm.get("title", "")),
        ("Описание в поиске", fm.get("description", "")),
    ]
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in info:
        cells = t.add_row().cells
        cells[0].text = k; shade(cells[0], "F2F2F2"); cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = str(v or "—")
    for row in t.rows:
        row.cells[0].width = Cm(5); row.cells[1].width = Cm(11.6)
    doc.add_paragraph()

    words = render_body(doc, body)
    info_p = doc.paragraphs[1]
    info_p.add_run(f" Объём: около {words} слов.").italic = True

    # Шаблон python-docx пишет <w:zoom> без обязательного percent — по схеме это ошибка
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{fm.get('slug', slug)}.docx"
    doc.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("slugs", nargs="+")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--notes", type=Path, default=None)
    a = ap.parse_args()
    notes = yaml.safe_load(a.notes.read_text(encoding="utf-8")) if a.notes else {}
    out = a.out or VAULT / "clients" / "arcars" / "reports" / f"drafts-{dt.date.today():%Y-%m-%d}"
    for s in a.slugs:
        print(build(s, out, (notes or {}).get(s)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
