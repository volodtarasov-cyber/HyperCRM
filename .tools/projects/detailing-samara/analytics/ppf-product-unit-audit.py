#!/usr/bin/env python3
"""Audit how Levsha PPF orders can be classified from the raw 1C export."""

from __future__ import annotations

import argparse
import csv
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[4]
DEFAULT_SOURCE = Path.home() / "Downloads/24_26_клиенты_услуги по месяцам.xlsx"
DEFAULT_OUTPUT = ROOT / ".tools/snapshots/ppf-product-unit-audit.csv"
DEFAULT_PREMIUM_OUTPUT = ROOT / ".tools/snapshots/ppf-premium-orders-review.csv"
REVENUE_COLUMN = 9

GENERIC_WORK = "Антигравийная пленка на кузов"
FULL_BODY_WORK = "Антигравийная пленка (Полная оклейка ТС)"
FILM_MATERIAL = "Материал Пленка"
SUSPICIOUS_CLIENT_TOKENS = (
    "[",
    "стеклоомыватель",
    "щетки стеклоочистителя",
    "спрей для",
)


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_revenue(value: object) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def parse_order_header(value: str) -> tuple[str, datetime | None]:
    match = re.search(r"Заказ-наряд\s+(\S+)\s+от\s+(\d{2}\.\d{2}\.\d{4})", value)
    if not match:
        return value, None
    return match.group(1), datetime.strptime(match.group(2), "%d.%m.%Y")


def classify(order: dict[str, object]) -> str:
    generic = float(order["generic_work"])
    full_body = float(order["full_body_work"])
    material = float(order["film_material"])

    if full_body > 0 and generic > 0:
        return "mixed_full_and_generic"
    if full_body > 0:
        return "explicit_full_body"
    if generic > 0 and material > 0:
        return "generic_work_plus_material"
    if generic > 0:
        return "generic_work_only"
    if material > 0:
        return "material_only"
    return "not_ppf"


def client_data_quality(client: object) -> str:
    value = normalize(client).lower()
    if any(token in value for token in SUSPICIOUS_CLIENT_TOKENS):
        return "suspect_product_name"
    return "ok"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--premium-output", type=Path, default=DEFAULT_PREMIUM_OUTPUT
    )
    parser.add_argument("--start-date", default="2025-01-01")
    args = parser.parse_args()

    cutoff = datetime.strptime(args.start_date, "%Y-%m-%d")
    workbook = load_workbook(args.source, read_only=True, data_only=True)
    worksheet = workbook.active

    orders: list[dict[str, object]] = []
    current_client = ""
    current_order: dict[str, object] | None = None

    for row in worksheet.iter_rows(min_row=15, values_only=True):
        first = normalize(row[0])
        if not first:
            continue

        if first.startswith("Заказ-наряд"):
            if current_order is not None:
                orders.append(current_order)
            order_id, order_date = parse_order_header(first)
            current_order = {
                "client": current_client,
                "order_id": order_id,
                "date": order_date,
                "total_order_revenue": parse_revenue(row[REVENUE_COLUMN]),
                "generic_work": 0.0,
                "full_body_work": 0.0,
                "film_material": 0.0,
                "services": [],
            }
            continue

        if first.startswith(", "):
            if current_order is None:
                continue
            service = normalize(first[2:])
            revenue = parse_revenue(row[REVENUE_COLUMN])
            current_order["services"].append((service, revenue))
            if service == GENERIC_WORK:
                current_order["generic_work"] += revenue
            elif service == FULL_BODY_WORK:
                current_order["full_body_work"] += revenue
            elif service == FILM_MATERIAL:
                current_order["film_material"] += revenue
            continue

        if current_order is not None:
            orders.append(current_order)
            current_order = None
        current_client = first

    if current_order is not None:
        orders.append(current_order)

    audited = []
    for order in orders:
        order_date = order["date"]
        if not isinstance(order_date, datetime) or order_date < cutoff:
            continue

        film_check = (
            float(order["generic_work"])
            + float(order["full_body_work"])
            + float(order["film_material"])
        )
        if film_check <= 0:
            continue

        order["film_check"] = film_check
        order["classification"] = classify(order)
        order["other_revenue"] = max(
            0.0, float(order["total_order_revenue"]) - film_check
        )
        audited.append(order)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "client",
                "order_id",
                "date",
                "classification",
                "total_order_revenue",
                "film_check",
                "generic_work",
                "full_body_work",
                "film_material",
                "other_revenue",
                "service_names",
            ]
        )
        for order in sorted(audited, key=lambda item: item["date"]):
            writer.writerow(
                [
                    order["client"],
                    order["order_id"],
                    order["date"].strftime("%Y-%m-%d"),
                    order["classification"],
                    order["total_order_revenue"],
                    order["film_check"],
                    order["generic_work"],
                    order["full_body_work"],
                    order["film_material"],
                    order["other_revenue"],
                    " | ".join(name for name, _ in order["services"]),
                ]
            )

    classifications = Counter(str(order["classification"]) for order in audited)
    premium = [order for order in audited if float(order["film_check"]) >= 300_000]
    premium_classes = Counter(str(order["classification"]) for order in premium)
    explicit = [
        order
        for order in audited
        if order["classification"] in {"explicit_full_body", "mixed_full_and_generic"}
    ]

    args.premium_output.parent.mkdir(parents=True, exist_ok=True)
    with args.premium_output.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "date",
                "order_id",
                "client",
                "client_data_quality",
                "classification",
                "full_body_status",
                "film_check",
                "total_order_revenue",
                "manual_product_unit",
                "car",
                "manager",
                "source",
                "notes",
            ]
        )
        for order in sorted(
            premium, key=lambda item: float(item["film_check"]), reverse=True
        ):
            is_explicit = order["classification"] in {
                "explicit_full_body",
                "mixed_full_and_generic",
            }
            writer.writerow(
                [
                    order["date"].strftime("%Y-%m-%d"),
                    order["order_id"],
                    order["client"],
                    client_data_quality(order["client"]),
                    order["classification"],
                    "explicit" if is_explicit else "needs_manual_review",
                    order["film_check"],
                    order["total_order_revenue"],
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )

    print(f"PPF orders since {args.start_date}: {len(audited)}")
    print(f"PPF revenue: {sum(float(order['film_check']) for order in audited):,.0f}")
    print("Classifications:")
    for label, count in classifications.most_common():
        print(f"  {label}: {count}")
    print(f"Explicit full-body orders: {len(explicit)}")
    print(
        "Explicit full-body revenue: "
        f"{sum(float(order['film_check']) for order in explicit):,.0f}"
    )
    print(f"Premium orders 300k+: {len(premium)}")
    print("Premium classifications:")
    for label, count in premium_classes.most_common():
        print(f"  {label}: {count}")
    print(f"CSV: {args.output}")
    print(f"Premium review CSV: {args.premium_output}")


if __name__ == "__main__":
    main()
