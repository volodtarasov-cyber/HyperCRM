#!/usr/bin/env python3
"""Локальное обновление аналитики Левши из ручных CSV-выгрузок."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import sqlite3
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover - Python 3.9+ в рабочей среде
    ZoneInfo = None


SUPPORTED_SUFFIXES = {".csv", ".xlsx"}
MIN_ATTRIBUTION_RATE = 0.8
PHONE_MATCH_LOOKBACK_DAYS = 45
STRICT_MATCH_MAX_DAY_DISTANCE = 3
LEAD_DEAL_MATCH_TOLERANCE_SECONDS = 10
PERIOD_RE = re.compile(r"(?P<start>20\d{2}-\d{2}-\d{2})_(?P<end>20\d{2}-\d{2}-\d{2})")
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d.%m.%Y",
    "%d.%m.%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%d/%m/%Y",
    "%Y/%m/%d",
)


def now_local() -> datetime:
    if ZoneInfo is not None:
        return datetime.now(ZoneInfo("Europe/Samara"))
    return datetime.now()


def iso_now() -> str:
    return now_local().isoformat(timespec="seconds")


def display_now() -> str:
    return now_local().strftime("%d.%m.%Y %H:%M")


def normalize_text(value: object) -> str:
    return str(value or "").replace("\ufeff", "").replace("\u00a0", " ").strip()


def normalize_key(value: object) -> str:
    return re.sub(r"\s+", " ", normalize_text(value)).casefold()


def normalize_order_no(value: object) -> str:
    text = normalize_text(value).replace(" ", "")
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text.upper()


def normalize_phone(value: object) -> str:
    digits = re.sub(r"\D", "", normalize_text(value))
    return digits[-10:] if len(digits) >= 10 else ""


def extract_phone_numbers(value: object) -> List[str]:
    text = normalize_text(value)
    if not text:
        return []
    chunks = re.split(r"[,;/\n]+", text)
    return sorted({phone for chunk in chunks if (phone := normalize_phone(chunk))})


def contact_phones(row: Dict[str, str]) -> List[str]:
    phones = set()
    for key, value in row.items():
        normalized_key = normalize_key(key)
        if normalized_key.startswith("контакт:") and "телефон" in normalized_key:
            phones.update(extract_phone_numbers(value))
    return sorted(phones)


def lead_phones(row: Dict[str, str]) -> List[str]:
    phones = set(contact_phones(row))
    for name in ("Рабочий телефон", "Мобильный телефон", "Телефон"):
        phones.update(extract_phone_numbers(field(row, name)))
    return sorted(phones)


def order_match_key(value: object) -> str:
    text = normalize_order_no(value).replace("Ё", "Е")
    if not text or text in {"Б/З", "Б/Н", "БЕЗЗН", "БЕЗЗАКАЗ-НАРЯДА"}:
        return ""
    letters = "".join(re.findall(r"[A-ZА-Я]", text))
    digits = "".join(re.findall(r"\d", text)).lstrip("0") or "0"
    if letters.startswith("ЛЭЛВ") or letters.startswith("ЛЭ"):
        return f"ЛЭ:{digits}"
    if letters.startswith("ЭЛЛВ") or letters.startswith("ЭЛ"):
        return f"ЭЛ:{digits}"
    if letters in {"", "ЛВ"}:
        return f"NUM:{digits}"
    return f"RAW:{text}"


def parse_number(value: object) -> float:
    text = normalize_text(value)
    if not text:
        return 0.0
    text = text.replace("₽", "").replace("руб.", "").replace("руб", "")
    text = text.replace(" ", "").replace("\u00a0", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text in {"", "-", ".", "-."}:
        return 0.0
    return float(text)


def parse_integer(value: object) -> int:
    return int(round(parse_number(value)))


def parse_date(value: object) -> Optional[str]:
    text = normalize_text(value)
    if not text:
        return None
    if "T" in text:
        text = text.replace("T", " ").split("+")[0]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    for cut in (10, 19):
        candidate = text[:cut]
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(candidate, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def parse_datetime(value: object) -> Optional[str]:
    text = normalize_text(value)
    if not text:
        return None
    if "T" in text:
        text = text.replace("T", " ").split("+")[0]
    for fmt in ("%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text[:19], fmt)
            return parsed.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    date = parse_date(text)
    return f"{date} 00:00:00" if date else None


def file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    raw = path.read_bytes()
    decoded: Optional[str] = None
    errors: List[str] = []
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            decoded = raw.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
    if decoded is None:
        raise ValueError("Не удалось определить кодировку CSV: " + "; ".join(errors))

    sample = decoded[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"

    reader = csv.DictReader(io.StringIO(decoded), delimiter=delimiter)
    if reader.fieldnames is None:
        raise ValueError("В файле не найдена строка заголовков")
    headers = [normalize_text(header) for header in reader.fieldnames]
    rows: List[Dict[str, str]] = []
    for source_row in reader:
        row = {normalize_text(key): normalize_text(value) for key, value in source_row.items() if key is not None}
        if any(row.values()):
            rows.append(row)
    return headers, rows


def read_xlsx_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Для XLSX не найден модуль openpyxl. Сохраните выгрузку в CSV UTF-8 с разделителем ';'"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        iterator = worksheet.iter_rows(values_only=True)
        header_values: Optional[Tuple[object, ...]] = None
        for candidate in iterator:
            if any(normalize_text(value) for value in candidate):
                header_values = candidate
                break
        if header_values is None:
            raise ValueError("В XLSX не найдена строка заголовков")
        headers = [normalize_text(value) for value in header_values]
        if not any(headers):
            raise ValueError("В XLSX строка заголовков пустая")

        rows: List[Dict[str, str]] = []
        for source_row in iterator:
            row = {
                header: normalize_text(source_row[index] if index < len(source_row) else "")
                for index, header in enumerate(headers)
                if header
            }
            if any(row.values()):
                rows.append(row)
        return headers, rows
    finally:
        workbook.close()


def read_table_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    if path.suffix.lower() == ".csv":
        return read_csv_rows(path)
    if path.suffix.lower() == ".xlsx":
        return read_xlsx_rows(path)
    raise ValueError("Поддерживаются CSV и XLSX")


def field(row: Dict[str, str], name: str) -> str:
    return normalize_text(row.get(name, ""))


def find_header(headers: Sequence[str], aliases: Sequence[str], required: bool = False) -> Optional[str]:
    by_key = {normalize_key(header): header for header in headers}
    for alias in aliases:
        if normalize_key(alias) in by_key:
            return by_key[normalize_key(alias)]
    if required:
        raise ValueError(f"Не найдена обязательная колонка. Ожидалась одна из: {', '.join(aliases)}")
    return None


def filename_period(path: Path, dates: Iterable[str]) -> Tuple[str, str]:
    match = PERIOD_RE.search(path.stem)
    if match:
        start, end = match.group("start"), match.group("end")
        if start > end:
            raise ValueError("В имени файла начало периода позже окончания")
        return start, end
    valid = sorted(date for date in dates if date)
    if not valid:
        raise ValueError("Не удалось определить период: добавьте даты в имя файла")
    return valid[0], valid[-1]


def archive_path(project: Path, source: str, path: Path, error: bool = False) -> Path:
    folder_name = "errors" if error else "archive"
    folder = project / folder_name / source
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / path.name
    if target.exists():
        stamp = now_local().strftime("%Y%m%d-%H%M%S")
        target = folder / f"{path.stem}_{stamp}{path.suffix}"
    return target


def move_processed(project: Path, source: str, path: Path, error: bool = False) -> None:
    target = archive_path(project, source, path, error=error)
    shutil.move(str(path), str(target))


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_manual_matches(path: Path) -> Dict[Tuple[str, str], str]:
    if not path.exists():
        return {}
    headers, rows = read_csv_rows(path)
    required = {"Номер заказ-наряда 1С", "ID сделки Б24"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Таблица ручных связок не содержит колонки: " + ", ".join(missing))

    matches: Dict[Tuple[str, str], str] = {}
    for index, row in enumerate(rows, start=2):
        order_no = normalize_order_no(field(row, "Номер заказ-наряда 1С"))
        deal_id = field(row, "ID сделки Б24")
        sale_date_raw = field(row, "Дата реализации 1С")
        sale_date = parse_date(sale_date_raw) if sale_date_raw else ""
        if not order_no or not deal_id:
            raise ValueError(f"Ручная связка, строка {index}: нужен заказ 1С и ID сделки Б24")
        if sale_date_raw and not sale_date:
            raise ValueError(f"Ручная связка, строка {index}: не распознана дата реализации")
        key = (order_no, sale_date)
        existing = matches.get(key)
        if existing and existing != deal_id:
            raise ValueError(f"Ручная связка, строка {index}: заказ указан с разными ID сделки")
        matches[key] = deal_id
    return matches


def init_database(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        PRAGMA journal_mode=DELETE;
        PRAGMA foreign_keys=ON;

        CREATE TABLE IF NOT EXISTS b24_deals (
            deal_id TEXT PRIMARY KEY,
            funnel TEXT,
            repeat_flag TEXT,
            segment TEXT NOT NULL,
            stage TEXT,
            previous_stage TEXT,
            success_flag INTEGER NOT NULL DEFAULT 0,
            booked_flag INTEGER NOT NULL DEFAULT 0,
            source_raw TEXT,
            source_normalized TEXT NOT NULL,
            source_extra TEXT,
            b24_amount REAL NOT NULL DEFAULT 0,
            currency TEXT,
            is_closed TEXT,
            date_created TEXT NOT NULL,
            date_modified TEXT,
            date_start TEXT,
            planned_close_date TEXT,
            event_date TEXT,
            stage_changed_at TEXT,
            utm_source TEXT,
            utm_medium TEXT,
            utm_campaign TEXT,
            utm_content TEXT,
            utm_term TEXT,
            vehicle TEXT,
            order_no TEXT,
            appointment_at TEXT,
            contact_channel TEXT,
            entry_point TEXT,
            source_file TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            raw_json TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_b24_created ON b24_deals(date_created);
        CREATE INDEX IF NOT EXISTS idx_b24_modified ON b24_deals(date_modified);
        CREATE INDEX IF NOT EXISTS idx_b24_order ON b24_deals(order_no);

        CREATE TABLE IF NOT EXISTS b24_deal_phones (
            deal_id TEXT NOT NULL,
            phone TEXT NOT NULL,
            PRIMARY KEY(deal_id, phone),
            FOREIGN KEY(deal_id) REFERENCES b24_deals(deal_id) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_b24_phone ON b24_deal_phones(phone);

        CREATE TABLE IF NOT EXISTS b24_leads (
            lead_id TEXT PRIMARY KEY,
            title TEXT,
            stage TEXT NOT NULL,
            qualified_flag INTEGER NOT NULL DEFAULT 0,
            repeat_flag TEXT,
            segment TEXT NOT NULL,
            direction TEXT NOT NULL,
            source_raw TEXT,
            source_normalized TEXT NOT NULL,
            source_extra TEXT,
            date_created TEXT NOT NULL,
            created_at TEXT NOT NULL,
            date_modified TEXT,
            stage_changed_at TEXT,
            utm_source TEXT,
            utm_medium TEXT,
            utm_campaign TEXT,
            utm_content TEXT,
            utm_term TEXT,
            source_file TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            raw_json TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_b24_lead_created ON b24_leads(date_created);
        CREATE INDEX IF NOT EXISTS idx_b24_lead_stage ON b24_leads(stage);
        CREATE INDEX IF NOT EXISTS idx_b24_lead_direction ON b24_leads(direction);

        CREATE TABLE IF NOT EXISTS b24_lead_phones (
            lead_id TEXT NOT NULL,
            phone TEXT NOT NULL,
            PRIMARY KEY(lead_id, phone),
            FOREIGN KEY(lead_id) REFERENCES b24_leads(lead_id) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_b24_lead_phone ON b24_lead_phones(phone);

        CREATE TABLE IF NOT EXISTS onec_rows (
            row_id TEXT PRIMARY KEY,
            order_no TEXT NOT NULL,
            sale_date TEXT NOT NULL,
            revenue REAL NOT NULL DEFAULT 0,
            correction INTEGER NOT NULL DEFAULT 0,
            customer TEXT,
            phone TEXT,
            period_start TEXT NOT NULL,
            period_end TEXT NOT NULL,
            source_file TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            raw_json TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_onec_date ON onec_rows(sale_date);
        CREATE INDEX IF NOT EXISTS idx_onec_order ON onec_rows(order_no);

        CREATE TABLE IF NOT EXISTS onec_row_phones (
            row_id TEXT NOT NULL,
            phone TEXT NOT NULL,
            PRIMARY KEY(row_id, phone),
            FOREIGN KEY(row_id) REFERENCES onec_rows(row_id) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_onec_phone ON onec_row_phones(phone);

        CREATE TABLE IF NOT EXISTS ads_rows (
            row_id TEXT PRIMARY KEY,
            activity_date TEXT NOT NULL,
            platform TEXT NOT NULL,
            account TEXT,
            campaign TEXT,
            campaign_id TEXT,
            ad_group TEXT,
            ad_group_id TEXT,
            ad_id TEXT,
            ad_title TEXT,
            condition_type TEXT,
            keyword TEXT,
            search_query TEXT,
            impressions INTEGER NOT NULL DEFAULT 0,
            clicks INTEGER NOT NULL DEFAULT 0,
            spend REAL NOT NULL DEFAULT 0,
            leads REAL NOT NULL DEFAULT 0,
            period_start TEXT NOT NULL,
            period_end TEXT NOT NULL,
            source_file TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            raw_json TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_ads_date ON ads_rows(activity_date);
        CREATE INDEX IF NOT EXISTS idx_ads_platform ON ads_rows(platform);

        CREATE TABLE IF NOT EXISTS import_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source TEXT NOT NULL,
            file_name TEXT NOT NULL,
            file_hash TEXT NOT NULL,
            status TEXT NOT NULL,
            row_count INTEGER NOT NULL DEFAULT 0,
            period_start TEXT,
            period_end TEXT,
            imported_at TEXT NOT NULL,
            message TEXT
        );

        CREATE INDEX IF NOT EXISTS idx_import_hash ON import_log(file_hash, status);
        """
    )
    ads_columns = {row[1] for row in connection.execute("PRAGMA table_info(ads_rows)")}
    ads_optional_columns = {
        "campaign_id": "TEXT",
        "ad_group": "TEXT",
        "ad_group_id": "TEXT",
        "ad_id": "TEXT",
        "ad_title": "TEXT",
        "condition_type": "TEXT",
        "keyword": "TEXT",
        "search_query": "TEXT",
    }
    for column, sql_type in ads_optional_columns.items():
        if column not in ads_columns:
            connection.execute(f"ALTER TABLE ads_rows ADD COLUMN {column} {sql_type}")
    connection.execute(
        "CREATE INDEX IF NOT EXISTS idx_ads_platform_date_campaign "
        "ON ads_rows(platform, activity_date, campaign)"
    )
    connection.execute(
        "CREATE INDEX IF NOT EXISTS idx_ads_platform_date_keyword "
        "ON ads_rows(platform, activity_date, keyword)"
    )
    connection.execute("PRAGMA optimize")
    connection.commit()


def backfill_normalized_phones(connection: sqlite3.Connection) -> None:
    b24_prepared = []
    for deal_id, raw_json in connection.execute("SELECT deal_id, raw_json FROM b24_deals"):
        try:
            raw = json.loads(raw_json)
        except (TypeError, json.JSONDecodeError):
            continue
        b24_prepared.extend((deal_id, phone) for phone in contact_phones(raw))

    onec_prepared = []
    for row_id, phone in connection.execute("SELECT row_id, phone FROM onec_rows"):
        onec_prepared.extend((row_id, value) for value in extract_phone_numbers(phone))

    with connection:
        if b24_prepared:
            connection.executemany(
                "INSERT OR IGNORE INTO b24_deal_phones(deal_id, phone) VALUES (?, ?)",
                b24_prepared,
            )
        if onec_prepared:
            connection.executemany(
                "INSERT OR IGNORE INTO onec_row_phones(row_id, phone) VALUES (?, ?)",
                onec_prepared,
            )


def backfill_ad_campaign_ids(connection: sqlite3.Connection) -> None:
    aliases = {normalize_key(value) for value in ("campaign_id", "CampaignId", "CampaignID", "№ Кампании", "ID кампании")}
    updates = []
    for row_id, raw_json in connection.execute(
        "SELECT row_id, raw_json FROM ads_rows WHERE COALESCE(campaign_id, '') = ''"
    ):
        try:
            raw = json.loads(raw_json)
        except (TypeError, json.JSONDecodeError):
            continue
        campaign_id = next(
            (normalize_text(value) for key, value in raw.items() if normalize_key(key) in aliases),
            "",
        )
        if campaign_id:
            updates.append((campaign_id, row_id))
    if updates:
        with connection:
            connection.executemany("UPDATE ads_rows SET campaign_id = ? WHERE row_id = ?", updates)


def backfill_ad_titles(connection: sqlite3.Connection) -> None:
    aliases = {normalize_key(value) for value in ("ad_title", "Title", "Заголовок", "Название объявления")}
    updates = []
    for row_id, raw_json in connection.execute(
        "SELECT row_id, raw_json FROM ads_rows WHERE COALESCE(ad_title, '') = '' AND COALESCE(ad_id, '') != ''"
    ):
        try:
            raw = json.loads(raw_json)
        except (TypeError, json.JSONDecodeError):
            continue
        ad_title = next(
            (normalize_text(value) for key, value in raw.items() if normalize_key(key) in aliases),
            "",
        )
        if ad_title:
            updates.append((ad_title, row_id))
    if updates:
        with connection:
            connection.executemany("UPDATE ads_rows SET ad_title = ? WHERE row_id = ?", updates)


def stage_flags(stage: str, stage_map: dict) -> Tuple[int, int]:
    normalized = normalize_key(stage)
    success_exact = {normalize_key(value) for value in stage_map.get("success_exact", [])}
    success_contains = [normalize_key(value) for value in stage_map.get("success_contains", [])]
    booked_contains = [normalize_key(value) for value in stage_map.get("booked_contains", [])]
    success = normalized in success_exact or any(term and term in normalized for term in success_contains)
    booked = any(term and term in normalized for term in booked_contains)
    return int(success), int(booked)


def lead_qualified_flag(stage: str, lead_stage_map: dict) -> int:
    normalized = normalize_key(stage)
    qualified_exact = {
        normalize_key(value) for value in lead_stage_map.get("qualified_exact", [])
    }
    return int(normalized in qualified_exact)


def deal_outcome(stage: str, is_closed: object, outcome_rules: dict) -> str:
    normalized_stage = normalize_key(stage)
    won_exact = {normalize_key(value) for value in outcome_rules.get("won_exact", [])}
    if normalized_stage in won_exact:
        return "won"
    closed_values = {
        normalize_key(value) for value in outcome_rules.get("closed_values", [])
    }
    if normalize_key(is_closed) in closed_values:
        return "lost"
    return "open"


def qualification_flag(direction: str, stage: str, success: object, rules: dict) -> int:
    if int(success or 0):
        return 1
    normalized_direction = normalize_key(direction)
    normalized_stage = normalize_key(stage)
    excluded_directions = {
        normalize_key(value) for value in rules.get("exclude_direction_exact", [])
    }
    excluded_stages = {normalize_key(value) for value in rules.get("exclude_stage_exact", [])}
    excluded_stage_terms = [
        normalize_key(value) for value in rules.get("exclude_stage_contains", [])
    ]
    excluded = (
        normalized_direction in excluded_directions
        or normalized_stage in excluded_stages
        or any(term and term in normalized_stage for term in excluded_stage_terms)
    )
    return int(not excluded)


def normalize_source(row: Dict[str, str], source_rules: dict) -> str:
    source_fields = source_rules.get(
        "source_fields",
        ["UTM Source", "ДКТ: Источник", "Источник", "Канал обращения", "Точка входа обращения"],
    )
    medium_fields = source_rules.get("medium_fields", ["UTM Medium", "ДКТ: Канал"])
    candidates = [field(row, name) for name in source_fields]
    normalized_candidates = {normalize_key(value) for value in candidates if value}
    chosen = next((value for value in candidates if value), "")
    haystack = " ".join(sorted(normalized_candidates))
    medium = next(
        (normalize_key(field(row, name)) for name in medium_fields if field(row, name)),
        "",
    )

    for rule in source_rules.get("rules", []):
        equals_any = {normalize_key(value) for value in rule.get("source_equals_any", [])}
        contains_any = [normalize_key(value) for value in rule.get("source_contains_any", [])]
        allowed_medium = {normalize_key(value) for value in rule.get("medium_in", [])}
        if equals_any and not (normalized_candidates & equals_any):
            continue
        if contains_any and not any(term and term in haystack for term in contains_any):
            continue
        if allowed_medium and medium not in allowed_medium:
            continue
        target = normalize_text(rule.get("target"))
        if target:
            return target

    return chosen or normalize_text(source_rules.get("default")) or "Источник не указан"


def refresh_b24_sources(connection: sqlite3.Connection, source_rules: dict) -> None:
    prepared = []
    for deal_id, raw_json in connection.execute("SELECT deal_id, raw_json FROM b24_deals"):
        try:
            raw = json.loads(raw_json)
        except (TypeError, json.JSONDecodeError):
            continue
        prepared.append((normalize_source(raw, source_rules), deal_id))
    if prepared:
        with connection:
            connection.executemany(
                "UPDATE b24_deals SET source_normalized = ? WHERE deal_id = ?",
                prepared,
            )


def refresh_lead_sources(connection: sqlite3.Connection, source_rules: dict) -> None:
    prepared = []
    for lead_id, raw_json in connection.execute("SELECT lead_id, raw_json FROM b24_leads"):
        try:
            raw = json.loads(raw_json)
        except (TypeError, json.JSONDecodeError):
            continue
        prepared.append((normalize_source(raw, source_rules), lead_id))
    if prepared:
        with connection:
            connection.executemany(
                "UPDATE b24_leads SET source_normalized = ? WHERE lead_id = ?",
                prepared,
            )


def import_log(
    connection: sqlite3.Connection,
    source: str,
    path: Path,
    digest: str,
    status: str,
    row_count: int = 0,
    period: Tuple[Optional[str], Optional[str]] = (None, None),
    message: str = "",
) -> None:
    connection.execute(
        """
        INSERT INTO import_log(source, file_name, file_hash, status, row_count, period_start, period_end, imported_at, message)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (source, path.name, digest, status, row_count, period[0], period[1], iso_now(), message),
    )
    connection.commit()


def already_imported(connection: sqlite3.Connection, digest: str) -> bool:
    row = connection.execute(
        "SELECT 1 FROM import_log WHERE file_hash = ? AND status = 'success' LIMIT 1", (digest,)
    ).fetchone()
    return row is not None


def import_b24(
    connection: sqlite3.Connection,
    path: Path,
    stage_map: dict,
    source_rules: dict,
) -> Tuple[int, str, str]:
    headers, rows = read_table_rows(path)
    required = {"ID", "Повторная сделка", "Стадия сделки", "Дата создания", "Дата изменения"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Выгрузка Б24 не содержит колонки: " + ", ".join(missing))

    imported_at = iso_now()
    prepared = []
    prepared_phones = []
    dates = []
    for index, row in enumerate(rows, start=2):
        deal_id = field(row, "ID")
        created = parse_date(field(row, "Дата создания"))
        modified = parse_datetime(field(row, "Дата изменения"))
        if not deal_id:
            raise ValueError(f"Строка {index}: пустой ID сделки")
        if not created:
            raise ValueError(f"Строка {index}: не распознана дата создания")
        repeat_flag = field(row, "Повторная сделка").upper()
        segment = "repeat" if repeat_flag == "Y" else "new" if repeat_flag == "N" else "unknown"
        stage = field(row, "Стадия сделки")
        success, booked = stage_flags(stage, stage_map)
        order_no = normalize_order_no(field(row, "Номер заказ-наряда"))
        prepared.append(
            (
                deal_id,
                field(row, "Воронка"),
                repeat_flag,
                segment,
                stage,
                field(row, "Предыдущая стадия"),
                success,
                booked,
                field(row, "Источник"),
                normalize_source(row, source_rules),
                field(row, "Дополнительно об источнике"),
                parse_number(field(row, "Сумма")),
                field(row, "Валюта"),
                field(row, "Сделка закрыта"),
                created,
                modified,
                parse_date(field(row, "Дата начала")),
                parse_date(field(row, "Предполагаемая дата закрытия")),
                parse_date(field(row, "Дата события")),
                parse_datetime(field(row, "Дата изменения стадии")),
                field(row, "UTM Source"),
                field(row, "UTM Medium"),
                field(row, "UTM Campaign"),
                field(row, "UTM Content"),
                field(row, "UTM Term"),
                field(row, "АВТОМОБИЛЬ"),
                order_no,
                parse_datetime(field(row, "Дата и время записи")),
                field(row, "Канал обращения"),
                field(row, "Точка входа обращения"),
                path.name,
                imported_at,
                json.dumps(row, ensure_ascii=False, separators=(",", ":")),
            )
        )
        prepared_phones.extend((deal_id, phone) for phone in contact_phones(row))
        dates.append(created)

    with connection:
        connection.executemany(
            """
            INSERT INTO b24_deals(
                deal_id, funnel, repeat_flag, segment, stage, previous_stage, success_flag, booked_flag,
                source_raw, source_normalized, source_extra, b24_amount, currency, is_closed,
                date_created, date_modified, date_start, planned_close_date, event_date, stage_changed_at,
                utm_source, utm_medium, utm_campaign, utm_content, utm_term, vehicle, order_no,
                appointment_at, contact_channel, entry_point, source_file, imported_at, raw_json
            ) VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
            )
            ON CONFLICT(deal_id) DO UPDATE SET
                funnel=excluded.funnel,
                repeat_flag=excluded.repeat_flag,
                segment=excluded.segment,
                stage=excluded.stage,
                previous_stage=excluded.previous_stage,
                success_flag=excluded.success_flag,
                booked_flag=excluded.booked_flag,
                source_raw=excluded.source_raw,
                source_normalized=excluded.source_normalized,
                source_extra=excluded.source_extra,
                b24_amount=excluded.b24_amount,
                currency=excluded.currency,
                is_closed=excluded.is_closed,
                date_created=excluded.date_created,
                date_modified=excluded.date_modified,
                date_start=excluded.date_start,
                planned_close_date=excluded.planned_close_date,
                event_date=excluded.event_date,
                stage_changed_at=excluded.stage_changed_at,
                utm_source=excluded.utm_source,
                utm_medium=excluded.utm_medium,
                utm_campaign=excluded.utm_campaign,
                utm_content=excluded.utm_content,
                utm_term=excluded.utm_term,
                vehicle=excluded.vehicle,
                order_no=excluded.order_no,
                appointment_at=excluded.appointment_at,
                contact_channel=excluded.contact_channel,
                entry_point=excluded.entry_point,
                source_file=excluded.source_file,
                imported_at=excluded.imported_at,
                raw_json=excluded.raw_json
            WHERE b24_deals.date_modified IS NULL
               OR excluded.date_modified IS NULL
               OR excluded.date_modified >= b24_deals.date_modified
            """,
            prepared,
        )
        if prepared_phones:
            connection.executemany(
                "INSERT OR IGNORE INTO b24_deal_phones(deal_id, phone) VALUES (?, ?)",
                prepared_phones,
            )
    return len(prepared), min(dates), max(dates)


def import_leads(
    connection: sqlite3.Connection,
    path: Path,
    lead_stage_map: dict,
    source_rules: dict,
) -> Tuple[int, str, str]:
    headers, rows = read_table_rows(path)
    required = {"ID", "Стадия", "Дата создания", "Дата изменения", "Повторный лид"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Выгрузка лидов Б24 не содержит колонки: " + ", ".join(missing))

    imported_at = iso_now()
    prepared = []
    prepared_phones = []
    dates = []
    lead_ids = []
    for index, row in enumerate(rows, start=2):
        lead_id = field(row, "ID")
        created_at = parse_datetime(field(row, "Дата создания"))
        created = parse_date(field(row, "Дата создания"))
        modified = parse_datetime(field(row, "Дата изменения"))
        if not lead_id:
            raise ValueError(f"Строка {index}: пустой ID лида")
        if not created or not created_at:
            raise ValueError(f"Строка {index}: не распознана дата создания лида")
        repeat_flag = field(row, "Повторный лид").upper()
        segment = "repeat" if repeat_flag == "Y" else "new" if repeat_flag == "N" else "unknown"
        stage = field(row, "Стадия") or "Статус не указан"
        direction = field(row, "Причина обращения") or "Направление не указано"
        prepared.append(
            (
                lead_id,
                field(row, "Название лида"),
                stage,
                lead_qualified_flag(stage, lead_stage_map),
                repeat_flag,
                segment,
                direction,
                field(row, "Источник"),
                normalize_source(row, source_rules),
                field(row, "Дополнительно об источнике"),
                created,
                created_at,
                modified,
                parse_datetime(field(row, "Дата изменения стадии")),
                field(row, "UTM Source"),
                field(row, "UTM Medium"),
                field(row, "UTM Campaign"),
                field(row, "UTM Content"),
                field(row, "UTM Term"),
                path.name,
                imported_at,
                json.dumps(row, ensure_ascii=False, separators=(",", ":")),
            )
        )
        prepared_phones.extend((lead_id, phone) for phone in lead_phones(row))
        lead_ids.append((lead_id,))
        dates.append(created)

    with connection:
        connection.executemany(
            """
            INSERT INTO b24_leads(
                lead_id, title, stage, qualified_flag, repeat_flag, segment, direction,
                source_raw, source_normalized, source_extra, date_created, created_at,
                date_modified, stage_changed_at, utm_source, utm_medium, utm_campaign,
                utm_content, utm_term, source_file, imported_at, raw_json
            ) VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
            )
            ON CONFLICT(lead_id) DO UPDATE SET
                title=excluded.title,
                stage=excluded.stage,
                qualified_flag=excluded.qualified_flag,
                repeat_flag=excluded.repeat_flag,
                segment=excluded.segment,
                direction=excluded.direction,
                source_raw=excluded.source_raw,
                source_normalized=excluded.source_normalized,
                source_extra=excluded.source_extra,
                date_created=excluded.date_created,
                created_at=excluded.created_at,
                date_modified=excluded.date_modified,
                stage_changed_at=excluded.stage_changed_at,
                utm_source=excluded.utm_source,
                utm_medium=excluded.utm_medium,
                utm_campaign=excluded.utm_campaign,
                utm_content=excluded.utm_content,
                utm_term=excluded.utm_term,
                source_file=excluded.source_file,
                imported_at=excluded.imported_at,
                raw_json=excluded.raw_json
            WHERE b24_leads.date_modified IS NULL
               OR excluded.date_modified IS NULL
               OR excluded.date_modified >= b24_leads.date_modified
            """,
            prepared,
        )
        connection.executemany("DELETE FROM b24_lead_phones WHERE lead_id = ?", lead_ids)
        if prepared_phones:
            connection.executemany(
                "INSERT OR IGNORE INTO b24_lead_phones(lead_id, phone) VALUES (?, ?)",
                prepared_phones,
            )
    return len(prepared), min(dates), max(dates)


def refresh_stage_flags(connection: sqlite3.Connection, stage_map: dict) -> None:
    rows = connection.execute("SELECT deal_id, stage FROM b24_deals").fetchall()
    prepared = [(stage_flags(stage or "", stage_map)[0], stage_flags(stage or "", stage_map)[1], deal_id) for deal_id, stage in rows]
    with connection:
        connection.executemany(
            "UPDATE b24_deals SET success_flag = ?, booked_flag = ? WHERE deal_id = ?", prepared
        )


def refresh_lead_stage_flags(connection: sqlite3.Connection, lead_stage_map: dict) -> None:
    rows = connection.execute("SELECT lead_id, stage FROM b24_leads").fetchall()
    prepared = [
        (lead_qualified_flag(stage or "", lead_stage_map), lead_id)
        for lead_id, stage in rows
    ]
    if prepared:
        with connection:
            connection.executemany(
                "UPDATE b24_leads SET qualified_flag = ? WHERE lead_id = ?", prepared
            )


def read_onec_pivot_orders(path: Path, correction_terms: Sequence[str]) -> Optional[List[dict]]:
    if path.suffix.lower() != ".xlsx":
        return None
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Для XLSX не найден модуль openpyxl. Сохраните выгрузку в CSV UTF-8 с разделителем ';'"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        report_title = normalize_key(worksheet.cell(1, 1).value)
        customer_header = normalize_key(worksheet.cell(3, 1).value)
        order_header = normalize_key(worksheet.cell(4, 1).value)
        if not (
            "продажи по заказам" in report_title
            and "покупатель" in customer_header
            and "заказ покупателя" in order_header
        ):
            return None

        date_columns: List[Tuple[int, str]] = []
        for column in range(1, worksheet.max_column + 1):
            sale_date = parse_date(worksheet.cell(3, column).value)
            quantity_label = normalize_key(worksheet.cell(4, column).value)
            revenue_label = normalize_key(worksheet.cell(4, column + 1).value) if column < worksheet.max_column else ""
            if sale_date and "количество" in quantity_label and "выруч" in revenue_label:
                date_columns.append((column, sale_date))
        if not date_columns:
            raise ValueError(
                "В отчёте 1С «Продажи по заказам» не найдены дневные колонки количества и выручки"
            )

        order_re = re.compile(
            r"^Заказ-наряд\s+(.+?)\s+от\s+(\d{2}\.\d{2}\.\d{4})\s*$",
            re.IGNORECASE,
        )
        phone_re = re.compile(
            r"(?:\+?7|8)[\s()\-]*\d{3}[\s()\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}"
        )
        records: List[dict] = []
        current_customer = ""
        current_phone = ""
        current_order_records: List[dict] = []

        for source_row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            label = normalize_text(row[0] if row else "")
            if not label:
                continue
            order_match = order_re.match(label)
            if order_match:
                current_order_records = []
                order_no = normalize_order_no(order_match.group(1))
                for quantity_column, sale_date in date_columns:
                    quantity_index = quantity_column - 1
                    revenue_index = quantity_column
                    quantity_value = row[quantity_index] if quantity_index < len(row) else 0
                    revenue_value = row[revenue_index] if revenue_index < len(row) else 0
                    quantity = parse_number(quantity_value)
                    revenue = parse_number(revenue_value)
                    if quantity == 0 and revenue == 0:
                        continue
                    record = {
                        "order_no": order_no,
                        "sale_date": sale_date,
                        "revenue": revenue,
                        "correction": 0,
                        "customer": current_customer,
                        "phone": current_phone,
                        "phones": set(extract_phone_numbers(current_phone)),
                        "items": [],
                        "source_row": source_row_index,
                        "duplicate_rows": 1,
                    }
                    records.append(record)
                    current_order_records.append(record)
                continue

            if label.startswith(",") and current_order_records:
                item = label.lstrip(", ").strip()
                for record in current_order_records:
                    record["items"].append(item)
                continue

            current_order_records = []
            phone_match = phone_re.search(label)
            current_phone = normalize_text(phone_match.group(0)) if phone_match else ""
            current_customer = normalize_text(phone_re.sub("", label).strip(" ,"))

        if not records:
            raise ValueError("В отчёте 1С не найдены строки заказ-нарядов с дневными значениями")

        deduplicated: Dict[Tuple[str, str], dict] = {}
        for record in records:
            record["correction"] = int(
                any(
                    term and term in normalize_key(text)
                    for term in correction_terms
                    for text in [record["order_no"], *record["items"]]
                )
            )
            key = (record["order_no"], record["sale_date"])
            existing = deduplicated.get(key)
            if existing is None:
                deduplicated[key] = record
                continue
            if (
                existing["revenue"] != record["revenue"]
                or existing["correction"] != record["correction"]
            ):
                raise ValueError(
                    f"Заказ-наряд {record['order_no']} за {record['sale_date']} повторяется с разными значениями"
                )
            existing["duplicate_rows"] += 1
            existing["phones"].update(record["phones"])
            if not existing["phone"] and record["phone"]:
                existing["phone"] = record["phone"]
                existing["customer"] = record["customer"]

        parsed = []
        for record in deduplicated.values():
            raw = {
                "Заказ-наряд": record["order_no"],
                "Дата реализации": record["sale_date"],
                "Выручка": record["revenue"],
                "Покупатель": record["customer"],
                "Телефон": record["phone"],
                "Телефоны (нормализованные)": sorted(record["phones"]),
                "Услуги": record["items"],
                "Строка источника": record["source_row"],
                "Повторов строки в отчёте": record["duplicate_rows"],
            }
            parsed.append(
                {
                    "order_no": record["order_no"],
                    "sale_date": record["sale_date"],
                    "revenue": record["revenue"],
                    "correction": record["correction"],
                    "customer": record["customer"],
                    "phone": record["phone"],
                    "phones": sorted(record["phones"]),
                    "raw_json": json.dumps(raw, ensure_ascii=False, separators=(",", ":")),
                }
            )
        return parsed
    finally:
        workbook.close()


def import_onec(connection: sqlite3.Connection, path: Path, mappings: dict, stage_map: dict) -> Tuple[int, str, str]:
    correction_terms = [normalize_key(value) for value in stage_map.get("correction_contains", ["коррекц"])]
    parsed = read_onec_pivot_orders(path, correction_terms)
    if parsed is None:
        headers, rows = read_table_rows(path)
        cfg = mappings["onec"]
        order_header = find_header(headers, cfg["order_no"], required=True)
        date_header = find_header(headers, cfg["sale_date"], required=True)
        revenue_header = find_header(headers, cfg["revenue"], required=True)
        correction_header = find_header(headers, cfg["correction"], required=False)
        customer_header = find_header(headers, cfg["customer"], required=False)
        phone_header = find_header(headers, cfg["phone"], required=False)

        parsed = []
        for index, row in enumerate(rows, start=2):
            order_no = normalize_order_no(row.get(order_header, ""))
            sale_date = parse_date(row.get(date_header, ""))
            if not order_no:
                raise ValueError(f"Строка {index}: пустой номер заказ-наряда")
            if not sale_date:
                raise ValueError(f"Строка {index}: не распознана дата реализации")
            correction_raw = normalize_key(row.get(correction_header, "")) if correction_header else ""
            correction = int(
                correction_raw in {"y", "yes", "да", "1", "true"}
                or any(term and term in correction_raw for term in correction_terms)
            )
            parsed.append(
                {
                    "order_no": order_no,
                    "sale_date": sale_date,
                    "revenue": parse_number(row.get(revenue_header, "")),
                    "correction": correction,
                    "customer": normalize_text(row.get(customer_header, "")) if customer_header else "",
                    "phone": normalize_text(row.get(phone_header, "")) if phone_header else "",
                    "phones": extract_phone_numbers(row.get(phone_header, "")) if phone_header else [],
                    "raw_json": json.dumps(row, ensure_ascii=False, separators=(",", ":")),
                }
            )

    dates = [row["sale_date"] for row in parsed]
    imported_at = iso_now()

    period_start, period_end = filename_period(path, dates)
    for row in parsed:
        if not (period_start <= row["sale_date"] <= period_end):
            raise ValueError(
                f"Дата {row['sale_date']} выходит за период {period_start}—{period_end}, указанный в имени файла"
            )

    with connection:
        connection.execute(
            "DELETE FROM onec_rows WHERE sale_date BETWEEN ? AND ?", (period_start, period_end)
        )
        prepared = []
        prepared_phones = []
        for index, row in enumerate(parsed):
            row_id = hashlib.sha256(
                f"{path.name}|{index}|{row['order_no']}|{row['sale_date']}|{row['revenue']}|{row['raw_json']}".encode("utf-8")
            ).hexdigest()
            prepared.append(
                (
                    row_id,
                    row["order_no"],
                    row["sale_date"],
                    row["revenue"],
                    row["correction"],
                    row["customer"],
                    row["phone"],
                    period_start,
                    period_end,
                    path.name,
                    imported_at,
                    row["raw_json"],
                )
            )
            prepared_phones.extend((row_id, phone) for phone in row.get("phones", []))
        connection.executemany(
            """
            INSERT INTO onec_rows(
                row_id, order_no, sale_date, revenue, correction, customer, phone,
                period_start, period_end, source_file, imported_at, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            prepared,
        )
        if prepared_phones:
            connection.executemany(
                "INSERT OR IGNORE INTO onec_row_phones(row_id, phone) VALUES (?, ?)",
                prepared_phones,
            )
    return len(parsed), period_start, period_end


def infer_platform(path: Path) -> str:
    name = path.stem.casefold()
    if "yandex" in name or "direct" in name or "яндекс" in name or "директ" in name:
        return "Яндекс Директ"
    if re.search(r"(^|[_\- ])vk([_\- ]|$)", name) or "вконтакте" in name:
        return "VK"
    if "avito" in name or "авито" in name:
        return "Авито"
    return "Источник не указан"


def normalize_platform(value: str, fallback: str) -> str:
    haystack = normalize_key(value or fallback)
    if any(term in haystack for term in ("yandex", "яндекс", "direct", "директ")):
        return "Яндекс Директ"
    if "vk" in haystack or "вконтакте" in haystack:
        return "VK"
    if "avito" in haystack or "авито" in haystack:
        return "Авито"
    return normalize_text(value) or fallback


def infer_platform_from_rows(path: Path, headers: Sequence[str], rows: Sequence[Dict[str, str]]) -> str:
    inferred = infer_platform(path)
    if inferred != "Источник не указан":
        return inferred
    header_keys = {normalize_key(header) for header in headers}
    if {"№ кампании", "день", "расход, ₽", "конверсии"}.issubset(header_keys):
        return "Яндекс Директ"
    sample_parts = list(headers)
    for row in rows[:50]:
        sample_parts.extend(row.values())
    haystack = " ".join(normalize_key(value) for value in sample_parts)
    if "вконтакте" in haystack or "vk" in haystack:
        return "VK"
    if "яндекс" in haystack or "yandex" in haystack or "директ" in haystack:
        return "Яндекс Директ"
    if "авито" in haystack or "avito" in haystack:
        return "Авито"
    return inferred


def import_ads(connection: sqlite3.Connection, path: Path, mappings: dict) -> Tuple[int, str, str]:
    headers, rows = read_table_rows(path)
    cfg = mappings["ads"]
    date_header = find_header(headers, cfg["date"], required=True)
    platform_header = find_header(headers, cfg["platform"], required=False)
    account_header = find_header(headers, cfg["account"], required=False)
    campaign_header = find_header(headers, cfg["campaign"], required=False)
    campaign_id_header = find_header(headers, cfg["campaign_id"], required=False)
    ad_group_header = find_header(headers, cfg["ad_group"], required=False)
    ad_group_id_header = find_header(headers, cfg["ad_group_id"], required=False)
    ad_id_header = find_header(headers, cfg["ad_id"], required=False)
    ad_title_header = find_header(headers, cfg["ad_title"], required=False)
    condition_type_header = find_header(headers, cfg["condition_type"], required=False)
    keyword_header = find_header(headers, cfg["keyword"], required=False)
    search_query_header = find_header(headers, cfg["search_query"], required=False)
    impressions_header = find_header(headers, cfg["impressions"], required=False)
    clicks_header = find_header(headers, cfg["clicks"], required=False)
    spend_header = find_header(headers, cfg["spend"], required=True)
    leads_header = find_header(headers, cfg["leads"], required=False)

    fallback_platform = infer_platform_from_rows(path, headers, rows)
    if fallback_platform == "Яндекс Директ":
        detailed_headers = {
            "№ Группы": ad_group_id_header,
            "№ Объявления": ad_id_header,
            "Ключевая фраза": keyword_header,
            "Поисковый запрос": search_query_header,
        }
        missing_detailed = [label for label, header in detailed_headers.items() if not header]
        if missing_detailed:
            raise ValueError(
                "Для Яндекс Директа стандартом является детальная выгрузка. "
                "Не хватает колонок: " + ", ".join(missing_detailed)
            )
    parsed = []
    dates = []
    platforms = set()
    imported_at = iso_now()
    for index, row in enumerate(rows, start=2):
        activity_date = parse_date(row.get(date_header, ""))
        if not activity_date:
            row_values = {normalize_key(value) for value in row.values()}
            if row_values.intersection({"итого", "всего", "total"}):
                continue
            raise ValueError(f"Строка {index}: не распознана дата рекламы")
        platform = normalize_platform(row.get(platform_header, "") if platform_header else "", fallback_platform)
        dates.append(activity_date)
        platforms.add(platform)
        parsed.append(
            {
                "date": activity_date,
                "platform": platform,
                "account": normalize_text(row.get(account_header, "")) if account_header else "",
                "campaign": normalize_text(row.get(campaign_header, "")) if campaign_header else "",
                "campaign_id": normalize_text(row.get(campaign_id_header, "")) if campaign_id_header else "",
                "ad_group": normalize_text(row.get(ad_group_header, "")) if ad_group_header else "",
                "ad_group_id": normalize_text(row.get(ad_group_id_header, "")) if ad_group_id_header else "",
                "ad_id": normalize_text(row.get(ad_id_header, "")) if ad_id_header else "",
                "ad_title": normalize_text(row.get(ad_title_header, "")) if ad_title_header else "",
                "condition_type": normalize_text(row.get(condition_type_header, "")) if condition_type_header else "",
                "keyword": normalize_text(row.get(keyword_header, "")) if keyword_header else "",
                "search_query": normalize_text(row.get(search_query_header, "")) if search_query_header else "",
                "impressions": parse_integer(row.get(impressions_header, "")) if impressions_header else 0,
                "clicks": parse_integer(row.get(clicks_header, "")) if clicks_header else 0,
                "spend": parse_number(row.get(spend_header, "")),
                "leads": parse_number(row.get(leads_header, "")) if leads_header else 0,
                "raw_json": json.dumps(row, ensure_ascii=False, separators=(",", ":")),
            }
        )

    period_start, period_end = filename_period(path, dates)
    for row in parsed:
        if not (period_start <= row["date"] <= period_end):
            raise ValueError(
                f"Дата {row['date']} выходит за период {period_start}—{period_end}, указанный в имени файла"
            )

    with connection:
        placeholders = ",".join("?" for _ in platforms)
        connection.execute(
            f"DELETE FROM ads_rows WHERE activity_date BETWEEN ? AND ? AND platform IN ({placeholders})",
            (period_start, period_end, *sorted(platforms)),
        )
        prepared = []
        for index, row in enumerate(parsed):
            row_id = hashlib.sha256(
                f"{path.name}|{index}|{row['date']}|{row['platform']}|{row['account']}|{row['campaign_id']}|{row['campaign']}|{row['ad_group_id']}|{row['ad_id']}|{row['ad_title']}|{row['keyword']}|{row['search_query']}|{row['raw_json']}".encode("utf-8")
            ).hexdigest()
            prepared.append(
                (
                    row_id,
                    row["date"],
                    row["platform"],
                    row["account"],
                    row["campaign"],
                    row["campaign_id"],
                    row["ad_group"],
                    row["ad_group_id"],
                    row["ad_id"],
                    row["ad_title"],
                    row["condition_type"],
                    row["keyword"],
                    row["search_query"],
                    row["impressions"],
                    row["clicks"],
                    row["spend"],
                    row["leads"],
                    period_start,
                    period_end,
                    path.name,
                    imported_at,
                    row["raw_json"],
                )
            )
        connection.executemany(
            """
            INSERT INTO ads_rows(
                row_id, activity_date, platform, account, campaign, campaign_id,
                ad_group, ad_group_id, ad_id, ad_title, condition_type, keyword, search_query, impressions, clicks,
                spend, leads, period_start, period_end, source_file, imported_at, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            prepared,
        )
    return len(parsed), period_start, period_end


def process_file(
    connection: sqlite3.Connection,
    project: Path,
    source: str,
    path: Path,
    stage_map: dict,
    lead_stage_map: dict,
    mappings: dict,
    source_rules: dict,
    should_move: bool,
) -> Tuple[str, int]:
    digest = file_digest(path)
    if already_imported(connection, digest):
        print(f"  ↪ {path.name}: уже загружен, пропускаю")
        if should_move:
            move_processed(project, source, path)
        return "skipped", 0

    try:
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError("Поддерживаются CSV и XLSX")
        if source == "b24":
            count, start, end = import_b24(connection, path, stage_map, source_rules)
        elif source == "leads":
            count, start, end = import_leads(connection, path, lead_stage_map, source_rules)
        elif source == "1c":
            count, start, end = import_onec(connection, path, mappings, stage_map)
        elif source == "ads":
            count, start, end = import_ads(connection, path, mappings)
        else:  # pragma: no cover
            raise ValueError(f"Неизвестный источник: {source}")
        import_log(connection, source, path, digest, "success", count, (start, end), "")
        if should_move:
            move_processed(project, source, path)
        print(f"  ✓ {path.name}: {count} строк, период {start}—{end}")
        return "success", count
    except Exception as exc:
        import_log(connection, source, path, digest, "error", 0, (None, None), str(exc))
        if should_move:
            move_processed(project, source, path, error=True)
        print(f"  ✗ {path.name}: {exc}")
        return "error", 0


def aggregate_dashboard(
    connection: sqlite3.Connection,
    ad_sources: dict,
    manual_matches: Dict[Tuple[str, str], str],
    analytics_rules: dict,
    qualification_rules: dict,
    lead_stage_map: dict,
    outcome_rules: dict,
) -> dict:
    b24_rows = connection.execute(
        """
        SELECT deal_id, date_created, segment, source_normalized, success_flag, booked_flag,
               b24_amount, order_no, stage, date_modified, stage_changed_at, appointment_at,
               COALESCE(funnel, ''), raw_json, is_closed
        FROM b24_deals
        ORDER BY date_created, deal_id
        """
    ).fetchall()
    lead_rows = connection.execute(
        """
        SELECT lead_id, date_created, created_at, segment, direction, stage,
               qualified_flag, stage_changed_at, title, source_normalized
        FROM b24_leads
        ORDER BY date_created, lead_id
        """
    ).fetchall()
    lead_phones_by_id: Dict[str, set] = defaultdict(set)
    for lead_id, phone in connection.execute("SELECT lead_id, phone FROM b24_lead_phones"):
        lead_phones_by_id[lead_id].add(phone)
    onec_orders_raw = connection.execute(
        """
        SELECT order_no, MIN(sale_date) AS sale_date, SUM(revenue) AS revenue
        FROM onec_rows
        WHERE correction = 0
        GROUP BY order_no
        ORDER BY sale_date, order_no
        """
    ).fetchall()
    exclude_zero_revenue_orders = bool(
        analytics_rules.get("exclude_zero_revenue_orders", False)
    )
    excluded_zero_revenue_orders = sum(
        1 for _, _, revenue in onec_orders_raw if abs(float(revenue or 0)) < 0.005
    )
    onec_orders = [
        row
        for row in onec_orders_raw
        if not exclude_zero_revenue_orders or abs(float(row[2] or 0)) >= 0.005
    ]
    onec_phones_by_order: Dict[str, set] = defaultdict(set)
    for order_no, phone in connection.execute(
        """
        SELECT r.order_no, p.phone
        FROM onec_rows AS r
        JOIN onec_row_phones AS p ON p.row_id = r.row_id
        WHERE r.correction = 0
        """
    ):
        onec_phones_by_order[order_no].add(phone)
    ads_rows = connection.execute(
        """
        SELECT activity_date, platform, COALESCE(campaign_id, ''), COALESCE(campaign, ''),
               SUM(impressions), SUM(clicks), SUM(spend), SUM(leads)
        FROM ads_rows
        GROUP BY activity_date, platform, campaign_id, campaign
        ORDER BY activity_date, platform, campaign
        """
    ).fetchall()
    keyword_rows = connection.execute(
        """
        SELECT activity_date, platform, COALESCE(keyword, ''), COALESCE(condition_type, ''),
               SUM(impressions), SUM(clicks), SUM(spend), SUM(leads)
        FROM ads_rows
        WHERE platform = 'Яндекс Директ'
        GROUP BY activity_date, platform, keyword, condition_type
        ORDER BY activity_date, keyword, condition_type
        """
    ).fetchall()
    ad_detail_rows = connection.execute(
        """
        SELECT activity_date, platform, COALESCE(campaign_id, ''), COALESCE(campaign, ''),
               COALESCE(ad_group_id, ''), COALESCE(ad_group, ''),
               COALESCE(ad_id, ''), COALESCE(ad_title, ''),
               SUM(impressions), SUM(clicks), SUM(spend), SUM(leads)
        FROM ads_rows
        WHERE platform = 'Яндекс Директ'
          AND (COALESCE(ad_group_id, '') != '' OR COALESCE(ad_id, '') != '')
        GROUP BY activity_date, platform, campaign_id, campaign,
                 ad_group_id, ad_group, ad_id, ad_title
        ORDER BY activity_date, campaign, ad_group, ad_id
        """
    ).fetchall()

    deals_by_order: Dict[str, List[tuple]] = defaultdict(list)
    deals_by_phone: Dict[str, List[tuple]] = defaultdict(list)
    deal_phones_by_id: Dict[str, set] = defaultdict(set)
    deals_by_id = {row[0]: row for row in b24_rows}
    for row in sorted(b24_rows, key=lambda item: ((item[9] or ""), item[0]), reverse=True):
        match_key = order_match_key(row[7])
        if match_key:
            deals_by_order[match_key].append(row)
    order_to_deal = {
        match_key: rows[0]
        for match_key, rows in deals_by_order.items()
        if len({row[0] for row in rows}) == 1
    }
    for deal_id, phone in connection.execute("SELECT deal_id, phone FROM b24_deal_phones"):
        deal = deals_by_id.get(deal_id)
        if deal:
            deals_by_phone[phone].append(deal)
            deal_phones_by_id[deal_id].add(phone)

    revenue_by_deal: Dict[str, float] = defaultdict(float)
    onec_orders_by_deal: Dict[str, int] = defaultdict(int)
    daily_sales_map: Dict[Tuple[str, str, str], dict] = defaultdict(
        lambda: {"orders": 0, "revenue": 0.0}
    )
    matched_orders = 0
    matched_by_manual = 0
    matched_by_order = 0
    matched_by_phone = 0
    matched_by_phone_unique = 0
    matched_by_phone_stage_date = 0
    matched_by_phone_amount = 0
    matched_by_phone_nearest = 0
    unmatched_orders = 0
    ambiguous_orders = 0
    matched_revenue = 0.0
    for order_no, sale_date, revenue in onec_orders:
        match_key = order_match_key(order_no)
        manual_deal_id = manual_matches.get((normalize_order_no(order_no), sale_date))
        if manual_deal_id is None:
            manual_deal_id = manual_matches.get((normalize_order_no(order_no), ""))
        deal = None
        match_method = ""
        if manual_deal_id:
            deal = deals_by_id.get(manual_deal_id)
            if deal is None:
                raise ValueError(
                    f"Ручная связка для заказа {order_no}: сделка Б24 {manual_deal_id} не найдена"
                )
            match_method = "manual"
        if not deal:
            deal = order_to_deal.get(match_key)
            match_method = "order" if deal else ""
        candidates_by_id = {}
        if not deal:
            order_phones = onec_phones_by_order.get(order_no, set())
            phones = order_phones if len(order_phones) == 1 else set()
            order_candidates = deals_by_order.get(match_key, []) if match_key else []
            if len({row[0] for row in order_candidates}) > 1:
                candidates = [
                    row
                    for row in order_candidates
                    if any(row in deals_by_phone.get(phone, []) for phone in phones)
                ]
            else:
                sale_dt = datetime.strptime(sale_date, "%Y-%m-%d").date()
                candidates = []
                for phone in phones:
                    for row in deals_by_phone.get(phone, []):
                        created_dt = datetime.strptime(row[1], "%Y-%m-%d").date()
                        age_days = (sale_dt - created_dt).days
                        if 0 <= age_days <= PHONE_MATCH_LOOKBACK_DAYS:
                            candidates.append(row)
            candidates_by_id = {row[0]: row for row in candidates}
            if len(candidates_by_id) == 1:
                deal = next(iter(candidates_by_id.values()))
                match_method = "phone_unique"
            elif len(candidates_by_id) > 1:
                candidate_rows = list(candidates_by_id.values())
                sale_dt = datetime.strptime(sale_date, "%Y-%m-%d").date()

                def deal_date(value: object) -> Optional[date]:
                    parsed = parse_date(value)
                    return datetime.strptime(parsed, "%Y-%m-%d").date() if parsed else None

                successful = [row for row in candidate_rows if normalize_key(row[8]) == "сделка успешна"]
                same_stage_date = [row for row in successful if deal_date(row[10]) == sale_dt]
                same_appointment_date = [row for row in successful if deal_date(row[11]) == sale_dt]
                same_amount = [
                    row
                    for row in successful
                    if abs(float(row[6] or 0) - float(revenue or 0)) < 0.01
                ]
                if len(same_stage_date) == 1:
                    deal = same_stage_date[0]
                    match_method = "phone_stage_date"
                elif len(same_appointment_date) == 1:
                    deal = same_appointment_date[0]
                    match_method = "phone_stage_date"
                elif len(same_amount) == 1:
                    deal = same_amount[0]
                    match_method = "phone_amount"
                else:
                    dated_candidates = []
                    for row in successful:
                        stage_date = deal_date(row[10])
                        if stage_date is None:
                            continue
                        distance = abs((stage_date - sale_dt).days)
                        if distance <= STRICT_MATCH_MAX_DAY_DISTANCE:
                            dated_candidates.append((distance, row))
                    if dated_candidates:
                        nearest_distance = min(distance for distance, _ in dated_candidates)
                        nearest = [
                            row for distance, row in dated_candidates if distance == nearest_distance
                        ]
                        if len(nearest) == 1:
                            deal = nearest[0]
                            match_method = "phone_nearest"
        if deal:
            segment = deal[2]
            source = deal[3]
            revenue_by_deal[deal[0]] += float(revenue or 0)
            onec_orders_by_deal[deal[0]] += 1
            matched_orders += 1
            matched_by_manual += int(match_method == "manual")
            matched_by_order += int(match_method == "order")
            matched_by_phone += int(match_method.startswith("phone_"))
            matched_by_phone_unique += int(match_method == "phone_unique")
            matched_by_phone_stage_date += int(match_method == "phone_stage_date")
            matched_by_phone_amount += int(match_method == "phone_amount")
            matched_by_phone_nearest += int(match_method == "phone_nearest")
            matched_revenue += float(revenue or 0)
        elif (
            match_key
            and len({row[0] for row in deals_by_order.get(match_key, [])}) > 1
        ) or len(candidates_by_id) > 1:
            segment = "unknown"
            source = "Неоднозначное сопоставление"
            ambiguous_orders += 1
        else:
            segment = "unknown"
            source = "Не сопоставлено"
            unmatched_orders += 1
        key = (sale_date, segment, source)
        daily_sales_map[key]["orders"] += 1
        daily_sales_map[key]["revenue"] += float(revenue or 0)

    lead_to_deal: Dict[str, tuple] = {}
    lead_match_methods: Dict[str, int] = defaultdict(int)
    assigned_deal_ids = set()
    deal_match_rows_by_date: Dict[str, List[dict]] = defaultdict(list)
    for deal in b24_rows:
        try:
            raw_deal = json.loads(deal[13])
        except (TypeError, json.JSONDecodeError):
            raw_deal = {}
        deal_created_at = parse_datetime(field(raw_deal, "Дата создания"))
        if not deal_created_at:
            continue
        deal_match_rows_by_date[deal_created_at[:10]].append(
            {
                "deal": deal,
                "created_at": datetime.strptime(deal_created_at, "%Y-%m-%d %H:%M:%S"),
                "title": normalize_key(field(raw_deal, "Название сделки")),
                "phones": deal_phones_by_id.get(deal[0], set()),
            }
        )

    for lead in sorted(lead_rows, key=lambda row: ((row[7] or ""), row[0])):
        lead_id = lead[0]
        if not int(lead[6] or 0) or not lead[7]:
            continue
        lead_stage_changed = parse_datetime(lead[7])
        if not lead_stage_changed:
            continue
        lead_stage_changed_at = datetime.strptime(
            lead_stage_changed, "%Y-%m-%d %H:%M:%S"
        )
        candidates = [
            row
            for row in deal_match_rows_by_date.get(lead_stage_changed[:10], [])
            if row["deal"][0] not in assigned_deal_ids
            and abs((row["created_at"] - lead_stage_changed_at).total_seconds())
            <= LEAD_DEAL_MATCH_TOLERANCE_SECONDS
        ]
        phones = lead_phones_by_id.get(lead_id, set())
        title = normalize_key(lead[8])
        ranked_candidates = (
            (
                [
                    row
                    for row in candidates
                    if phones.intersection(row["phones"])
                    and title
                    and title == row["title"]
                ],
                "phone_title_time",
            ),
            (
                [row for row in candidates if phones.intersection(row["phones"])],
                "phone_time",
            ),
            (
                [row for row in candidates if title and title == row["title"]],
                "title_time",
            ),
            (candidates, "time_only"),
        )
        for matched_candidates, method in ranked_candidates:
            if len(matched_candidates) != 1:
                continue
            deal = matched_candidates[0]["deal"]
            lead_to_deal[lead_id] = deal
            assigned_deal_ids.add(deal[0])
            lead_match_methods[method] += 1
            break

    daily_funnel_map: Dict[Tuple[str, str, str], dict] = defaultdict(
        lambda: {"deals": 0, "success": 0, "booked": 0, "b24_amount": 0.0, "attributed_revenue": 0.0}
    )
    daily_direction_status_map: Dict[Tuple[str, str, str, str, str], dict] = defaultdict(
        lambda: {
            "leads": 0,
            "qualified": 0,
            "deals": 0,
            "won": 0,
            "open": 0,
            "lost": 0,
            "success": 0,
            "won_b24_amount": 0.0,
            "onec_orders": 0,
            "onec_revenue": 0.0,
        }
    )
    for deal in b24_rows:
        deal_id, created, segment, source, success, booked, amount = deal[:7]
        key = (created, segment, source)
        bucket = daily_funnel_map[key]
        bucket["deals"] += 1
        bucket["success"] += int(success or 0)
        bucket["booked"] += int(booked or 0)
        bucket["b24_amount"] += float(amount or 0)
        bucket["attributed_revenue"] += revenue_by_deal.get(deal_id, 0.0)

        if not lead_rows:
            direction = normalize_text(deal[12]) or "Направление не указано"
            status = normalize_text(deal[8]) or "Статус не указан"
            direction_bucket = daily_direction_status_map[
                (created, segment, source, direction, status)
            ]
            direction_bucket["leads"] += 1
            direction_bucket["qualified"] += qualification_flag(
                direction, status, success, qualification_rules
            )
            direction_bucket["deals"] += 1
            outcome = deal_outcome(deal[8], deal[14], outcome_rules)
            direction_bucket[outcome] += 1
            if outcome == "won":
                direction_bucket["won_b24_amount"] += float(deal[6] or 0)
            direction_bucket["success"] += int(success or booked or 0)
            direction_bucket["onec_orders"] += onec_orders_by_deal.get(deal_id, 0)
            direction_bucket["onec_revenue"] += revenue_by_deal.get(deal_id, 0.0)

    funnel_reporting_from = min((row[1] for row in b24_rows), default=None)
    if lead_rows:
        for lead in lead_rows:
            lead_id, created, _, lead_segment, direction, status, qualified = lead[:7]
            lead_source = normalize_text(lead[9]) or "Источник не указан"
            if funnel_reporting_from and created < funnel_reporting_from:
                continue
            matched_deal = lead_to_deal.get(lead_id)
            segment = matched_deal[2] if matched_deal else lead_segment
            direction_bucket = daily_direction_status_map[
                (
                    created,
                    segment,
                    lead_source,
                    normalize_text(direction) or "Направление не указано",
                    normalize_text(status) or "Статус не указан",
                )
            ]
            direction_bucket["leads"] += 1
            direction_bucket["qualified"] += int(qualified or 0)
            if matched_deal:
                direction_bucket["deals"] += 1
                outcome = deal_outcome(matched_deal[8], matched_deal[14], outcome_rules)
                direction_bucket[outcome] += 1
                if outcome == "won":
                    direction_bucket["won_b24_amount"] += float(matched_deal[6] or 0)
                direction_bucket["success"] += int(
                    matched_deal[4] or matched_deal[5] or 0
                )
                direction_bucket["onec_orders"] += onec_orders_by_deal.get(
                    matched_deal[0], 0
                )
                direction_bucket["onec_revenue"] += revenue_by_deal.get(
                    matched_deal[0], 0.0
                )

    daily_funnel = [
        {
            "date": date,
            "segment": segment,
            "source": source,
            **{key: round(value, 2) if isinstance(value, float) else value for key, value in values.items()},
        }
        for (date, segment, source), values in sorted(daily_funnel_map.items())
    ]
    daily_direction_status = [
        {
            "date": date,
            "segment": segment,
            "source": source,
            "direction": direction,
            "status": status,
            **{
                key: round(value, 2) if isinstance(value, float) else value
                for key, value in values.items()
            },
        }
        for (date, segment, source, direction, status), values in sorted(
            daily_direction_status_map.items()
        )
    ]
    daily_sales = [
        {
            "date": date,
            "segment": segment,
            "source": source,
            "orders": values["orders"],
            "revenue": round(values["revenue"], 2),
        }
        for (date, segment, source), values in sorted(daily_sales_map.items())
    ]
    daily_ads = []
    for row in ads_rows:
        platform = row[1]
        raw_spend = float(row[6] or 0)
        source_config = ad_sources.get(platform, {})
        vat_included = source_config.get("vat_included", True)
        vat_rate = float(source_config.get("vat_rate", 0) or 0)
        calculation_spend = raw_spend if vat_included else raw_spend * (1 + vat_rate)
        source_basis = source_config.get("spend_basis", "")
        if vat_included or not vat_rate:
            calculation_basis = source_basis
        else:
            calculation_basis = f"с НДС (+{vat_rate:.0%})"
        daily_ads.append(
            {
                "date": row[0],
                "platform": platform,
                "campaign_id": row[2],
                "campaign": row[3] or "Без названия кампании",
                "impressions": int(row[4] or 0),
                "clicks": int(row[5] or 0),
                "spend_raw": round(raw_spend, 2),
                "spend": round(calculation_spend, 4),
                "spend_basis": calculation_basis,
                "source_spend_basis": source_basis,
                "leads": round(float(row[7] or 0), 2),
                "conversions_available": bool(source_config.get("conversions_available", True)),
                "conversion_label": source_config.get("conversion_label", "Конверсии кабинета"),
                "conversion_definition": source_config.get("conversion_definition", ""),
            }
        )

    daily_keywords = []
    for row in keyword_rows:
        platform = row[1]
        raw_spend = float(row[6] or 0)
        source_config = ad_sources.get(platform, {})
        vat_included = source_config.get("vat_included", True)
        vat_rate = float(source_config.get("vat_rate", 0) or 0)
        calculation_spend = raw_spend if vat_included else raw_spend * (1 + vat_rate)
        daily_keywords.append(
            {
                "date": row[0],
                "platform": platform,
                "keyword": row[2],
                "condition_type": row[3],
                "impressions": int(row[4] or 0),
                "clicks": int(row[5] or 0),
                "spend": round(calculation_spend, 4),
                "leads": round(float(row[7] or 0), 2),
            }
        )

    daily_ad_details = []
    for row in ad_detail_rows:
        platform = row[1]
        raw_spend = float(row[10] or 0)
        source_config = ad_sources.get(platform, {})
        vat_included = source_config.get("vat_included", True)
        vat_rate = float(source_config.get("vat_rate", 0) or 0)
        calculation_spend = raw_spend if vat_included else raw_spend * (1 + vat_rate)
        daily_ad_details.append(
            {
                "date": row[0],
                "platform": platform,
                "campaign_id": row[2],
                "campaign": row[3] or "Без названия кампании",
                "ad_group_id": row[4],
                "ad_group": row[5] or "Без названия группы",
                "ad_id": row[6],
                "ad_title": row[7],
                "impressions": int(row[8] or 0),
                "clicks": int(row[9] or 0),
                "spend": round(calculation_spend, 4),
                "leads": round(float(row[11] or 0), 2),
            }
        )

    onec_row_count = connection.execute("SELECT COUNT(*) FROM onec_rows").fetchone()[0]
    ads_row_count = connection.execute("SELECT COUNT(*) FROM ads_rows").fetchone()[0]
    lead_row_count = len(lead_rows)
    success_without_order = connection.execute(
        "SELECT COUNT(*) FROM b24_deals WHERE success_flag = 1 AND COALESCE(order_no, '') = ''"
    ).fetchone()[0]
    last_import = connection.execute(
        "SELECT MAX(imported_at) FROM import_log WHERE status = 'success'"
    ).fetchone()[0]

    source_cutoffs = []
    if b24_rows:
        source_cutoffs.append(max(row[1] for row in b24_rows if row[1]))
    if onec_orders:
        source_cutoffs.append(max(row[1] for row in onec_orders if row[1]))
    if ads_rows:
        source_cutoffs.append(max(row[0] for row in ads_rows if row[0]))
    if lead_rows:
        source_cutoffs.append(max(row[1] for row in lead_rows if row[1]))
    reporting_through = min(source_cutoffs) if source_cutoffs else None

    source_parts = []
    if b24_rows:
        source_parts.append(f"Б24: {len(b24_rows)}")
    if lead_row_count:
        source_parts.append(f"лиды: {lead_row_count}")
    if onec_row_count:
        source_parts.append(f"1С: {onec_row_count}")
    if ads_row_count:
        source_parts.append(f"реклама: {ads_row_count}")

    attribution_match_rate = matched_orders / len(onec_orders) if onec_orders else 0.0
    deal_period_end = max((row[1] for row in b24_rows), default=None)
    eligible_qualified_leads = sum(
        1
        for row in lead_rows
        if int(row[6] or 0)
        and row[7]
        and funnel_reporting_from
        and deal_period_end
        and funnel_reporting_from <= row[7][:10] <= deal_period_end
    )
    lead_deal_match_rate = (
        len(lead_to_deal) / eligible_qualified_leads if eligible_qualified_leads else 0.0
    )
    actual_lead_funnel = bool(lead_rows)

    return {
        "meta": {
            "generated_at": iso_now(),
            "generated_at_display": display_now(),
            "source_caption": " · ".join(source_parts) if source_parts else "Данные ещё не загружены",
            "revenue_mode": "onec" if onec_orders else "b24",
            "reporting_through": reporting_through,
            "attribution_match_rate": round(attribution_match_rate, 4),
            "attribution_ready": bool(onec_orders and attribution_match_rate >= MIN_ATTRIBUTION_RATE),
            "roi_margin": round(float(analytics_rules.get("roi_margin", 0.3)), 4),
            "exclude_zero_revenue_orders": exclude_zero_revenue_orders,
            "qualification_label": normalize_text(
                lead_stage_map.get("label")
                if actual_lead_funnel
                else qualification_rules.get("label", "Квал — прокси по текущей стадии Б24")
            ),
            "qualification_description": normalize_text(
                lead_stage_map.get("description")
                if actual_lead_funnel
                else qualification_rules.get("description", "")
            ),
            "deal_outcome_description": normalize_text(
                outcome_rules.get("description", "")
            ),
            "funnel_reporting_from": funnel_reporting_from,
            "funnel_uses_leads": actual_lead_funnel,
        },
        "daily_funnel": daily_funnel,
        "daily_direction_status": daily_direction_status,
        "daily_sales": daily_sales,
        "daily_ads": daily_ads,
        "daily_keywords": daily_keywords,
        "daily_ad_details": daily_ad_details,
        "quality": {
            "b24_deals": len(b24_rows),
            "b24_leads": lead_row_count,
            "eligible_qualified_leads": eligible_qualified_leads,
            "matched_leads_to_deals": len(lead_to_deal),
            "lead_deal_match_rate": round(lead_deal_match_rate, 4),
            "lead_match_phone_title_time": lead_match_methods["phone_title_time"],
            "lead_match_phone_time": lead_match_methods["phone_time"],
            "lead_match_title_time": lead_match_methods["title_time"],
            "lead_match_time_only": lead_match_methods["time_only"],
            "onec_rows": onec_row_count,
            "onec_orders": len(onec_orders),
            "onec_orders_raw": len(onec_orders_raw),
            "excluded_zero_revenue_orders": (
                excluded_zero_revenue_orders if exclude_zero_revenue_orders else 0
            ),
            "matched_orders": matched_orders,
            "matched_by_manual": matched_by_manual,
            "matched_by_order": matched_by_order,
            "matched_by_phone": matched_by_phone,
            "matched_by_phone_unique": matched_by_phone_unique,
            "matched_by_phone_stage_date": matched_by_phone_stage_date,
            "matched_by_phone_amount": matched_by_phone_amount,
            "matched_by_phone_nearest": matched_by_phone_nearest,
            "unmatched_orders": unmatched_orders,
            "ambiguous_orders": ambiguous_orders,
            "matched_revenue": round(matched_revenue, 2),
            "success_without_order": success_without_order,
            "ads_rows": ads_row_count,
            "last_import_at": last_import.replace("T", " ")[:16] if last_import else "—",
        },
    }


def write_dashboard_data(project: Path, payload: dict) -> None:
    target = project / "dashboard" / "data.js"
    target.parent.mkdir(parents=True, exist_ok=True)
    content = "window.LEVSHA_DATA = " + json.dumps(
        payload, ensure_ascii=False, separators=(",", ":")
    ) + ";\n"
    temp = target.with_suffix(".js.tmp")
    temp.write_text(content, encoding="utf-8")
    temp.replace(target)


def write_log(project: Path, lines: Sequence[str]) -> None:
    path = project / "logs" / "update.log"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        for line in lines:
            handle.write(f"{iso_now()} {line}\n")


def source_files(project: Path, source: str) -> List[Path]:
    folder = project / "inbox" / source
    folder.mkdir(parents=True, exist_ok=True)
    return sorted(path for path in folder.iterdir() if path.is_file() and not path.name.startswith("."))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Обновить локальную аналитику Левши")
    parser.add_argument("--project", type=Path, required=True, help="Папка levsha-dashboard")
    parser.add_argument("--seed-b24", type=Path, help="Одноразово загрузить CSV Б24 без перемещения файла")
    parser.add_argument(
        "--seed-leads", type=Path, help="Одноразово загрузить CSV лидов Б24 без перемещения файла"
    )
    parser.add_argument("--no-inbox", action="store_true", help="Не обрабатывать папки inbox")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    project = args.project.expanduser().resolve()
    if not project.exists():
        print(f"Не найдена папка проекта: {project}", file=sys.stderr)
        return 2

    stage_map = load_json(project / "config" / "stage_map.json")
    lead_stage_map = load_json(project / "config" / "lead_stage_map.json")
    outcome_rules = load_json(project / "config" / "deal_outcome_rules.json")
    mappings = load_json(project / "config" / "import_mappings.json")
    ad_sources = load_json(project / "config" / "ad_sources.json")
    analytics_rules = load_json(project / "config" / "analytics_rules.json")
    qualification_rules = load_json(project / "config" / "qualification_rules.json")
    source_rules = load_json(project / "config" / "source_rules.json")
    manual_matches = load_manual_matches(project / "config" / "manual_matches.csv")
    database_path = project / "data" / "levsha.sqlite3"
    database_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(database_path))
    init_database(connection)
    backfill_normalized_phones(connection)
    backfill_ad_campaign_ids(connection)
    backfill_ad_titles(connection)

    successes = 0
    errors = 0
    imported_rows = 0
    print("Обновление аналитики Левши")

    if args.seed_b24:
        seed = args.seed_b24.expanduser().resolve()
        print("\nПервичная загрузка Б24:")
        status, count = process_file(
            connection,
            project,
            "b24",
            seed,
            stage_map,
            lead_stage_map,
            mappings,
            source_rules,
            should_move=False,
        )
        successes += int(status == "success")
        errors += int(status == "error")
        imported_rows += count

    if args.seed_leads:
        seed = args.seed_leads.expanduser().resolve()
        print("\nПервичная загрузка лидов Б24:")
        status, count = process_file(
            connection,
            project,
            "leads",
            seed,
            stage_map,
            lead_stage_map,
            mappings,
            source_rules,
            should_move=False,
        )
        successes += int(status == "success")
        errors += int(status == "error")
        imported_rows += count

    if not args.no_inbox:
        for source, label in (
            ("b24", "сделки Битрикс24"),
            ("leads", "лиды Битрикс24"),
            ("1c", "1С"),
            ("ads", "реклама"),
        ):
            files = source_files(project, source)
            if not files:
                continue
            print(f"\n{label}:")
            for path in files:
                status, count = process_file(
                    connection,
                    project,
                    source,
                    path,
                    stage_map,
                    lead_stage_map,
                    mappings,
                    source_rules,
                    should_move=True,
                )
                successes += int(status == "success")
                errors += int(status == "error")
                imported_rows += count

    refresh_stage_flags(connection, stage_map)
    refresh_lead_stage_flags(connection, lead_stage_map)
    refresh_b24_sources(connection, source_rules)
    refresh_lead_sources(connection, source_rules)
    payload = aggregate_dashboard(
        connection,
        ad_sources,
        manual_matches,
        analytics_rules,
        qualification_rules,
        lead_stage_map,
        outcome_rules,
    )
    write_dashboard_data(project, payload)
    write_log(
        project,
        [
            f"successes={successes} errors={errors} imported_rows={imported_rows}",
            f"b24={payload['quality']['b24_deals']} leads={payload['quality']['b24_leads']} onec={payload['quality']['onec_rows']} ads={payload['quality']['ads_rows']}",
        ],
    )
    connection.close()

    print("\nИтог:")
    print(f"  Сделок Б24 в базе: {payload['quality']['b24_deals']}")
    print(f"  Лидов Б24 в базе: {payload['quality']['b24_leads']}")
    print(f"  Строк 1С: {payload['quality']['onec_rows']}")
    print(f"  Строк рекламы: {payload['quality']['ads_rows']}")
    print(f"  Дашборд: {project / 'dashboard' / 'index.html'}")
    if errors:
        print(f"  Ошибок файлов: {errors}. Файлы перемещены в errors/.")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
